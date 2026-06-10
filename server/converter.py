# 1. Merge PDFs
def merge_pdfs(file_paths, output_path):
    from pypdf import PdfWriter
    writer = PdfWriter()
    for path in file_paths:
        writer.append(path)
    with open(output_path, "wb") as f:
        writer.write(f)

# 2. Merge Word files  
def merge_word(file_paths, output_path):
    from docx import Document
    from docx.oxml.ns import qn
    import copy

    if not file_paths:
        return

    # Open first document as base
    base = Document(file_paths[0])
    
    # Process subsequent documents one by one to save memory
    for path in file_paths[1:]:
        doc = Document(path)
        for element in doc.element.body:
            # Append each element using XML copy approach
            if element.tag != qn('w:sectPr'):
                base.element.body.insert(-1, copy.deepcopy(element))
        # Explicitly delete to free memory
        del doc
        
    base.save(output_path)

# 3. PDF to Word
def pdf_to_word(input_path, output_path):
    from pdf2docx import Converter
    cv = Converter(input_path)
    cv.convert(output_path)
    cv.close()

# 4. Excel to PDF
def excel_to_pdf(input_path, output_path):
    import openpyxl
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

    wb = openpyxl.load_file(input_path) if hasattr(openpyxl, 'load_file') else openpyxl.load_workbook(input_path)
    sheet = wb.active
    data = []
    for row in sheet.iter_rows(values_only=True):
        data.append([str(cell) if cell is not None else "" for cell in row])

    doc = SimpleDocTemplate(output_path, pagesize=A4)
    elements = []
    if data:
        t = Table(data)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), '#cccccc'),
            ('GRID', (0, 0), (-1, -1), 1, '#000000'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
        ]))
        elements.append(t)
    doc.build(elements)

if __name__ == "__main__":
    import sys
    import json

    if len(sys.argv) < 2:
        sys.exit(1)

    command = sys.argv[1]

    if command == "merge_pdfs":
        # args: output_path, input_path1, input_path2, ...
        output_path = sys.argv[2]
        file_paths = sys.argv[3:]
        merge_pdfs(file_paths, output_path)

    elif command == "merge_word":
        # args: output_path, input_path1, input_path2, ...
        output_path = sys.argv[2]
        file_paths = sys.argv[3:]
        merge_word(file_paths, output_path)

    elif command == "pdf_to_word":
        # args: input_path, output_path
        input_path = sys.argv[2]
        output_path = sys.argv[3]
        pdf_to_word(input_path, output_path)

    elif command == "excel_to_pdf":
        # args: input_path, output_path
        input_path = sys.argv[2]
        output_path = sys.argv[3]
        excel_to_pdf(input_path, output_path)